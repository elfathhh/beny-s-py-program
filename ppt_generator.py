from pptx import Presentation

prs = Presentation()

slides_content = [
("Parallel Processing",""),
("Anggota Kelompok","Nama 1\nNama 2\nNama 3"),
("Pengertian","Parallel processing adalah metode komputasi di mana beberapa proses dijalankan secara bersamaan.\nBerbeda dengan sequential processing yang menjalankan tugas satu per satu."),
("Latar Belakang","Kebutuhan komputasi terus meningkat, terutama untuk data besar.\nSingle processor memiliki keterbatasan, sehingga dibutuhkan solusi berupa parallel processing."),
("Konsep Dasar","Task decomposition: membagi tugas besar menjadi bagian kecil.\nData & task parallelism: cara membagi pekerjaan.\nTask parallelism\nSinkronisasi diperlukan agar proses tetap terkoordinasi."),
("Arsitektur Parallel Processing","SISD: satu instruksi satu data.\nSIMD: satu instruksi banyak data.\nMISD: banyak instruksi satu data.\nMIMD: banyak instruksi banyak data (paling umum digunakan)."),
("Jenis Parallel Processing","Bit-level: operasi pada bit dilakukan bersamaan.\nInstruction-level: eksekusi beberapa instruksi sekaligus.\nData & task parallelism: pembagian data atau tugas."),
("Cara Kerja","Membagi tugas menjadi beberapa bagian kecil\nDiproses oleh beberapa processor\nHasil digabungkan kembali menjadi output akhir"),
("Contoh Implementasi","Multi-core CPU: banyak inti dalam satu processor.\nGPU: memproses banyak data sekaligus.\nCluster & cloud: menggabungkan banyak komputer."),
("Kelebihan","Mempercepat proses komputasi.\nLebih efisien dalam penggunaan waktu.\nMudah dikembangkan (scalable)."),
("Kekurangan","Lebih kompleks untuk dibuat dan diprogram.\nBiaya perangkat lebih tinggi.\nSulit dalam sinkronisasi dan debugging."),
("Contoh Kasus","Machine learning untuk training model.\nSimulasi ilmiah dan pengolahan big data."),
("Tools & Teknologi","OpenMP & MPI: untuk pemrograman paralel.\nCUDA: untuk GPU computing.\nHadoop & Spark: untuk big data processing."),
("Kesimpulan","Parallel processing sangat penting di era modern.\nMemberikan peningkatan performa yang signifikan.Namun membutuhkan pengelolaan yang baik."),
("Terima Kasih", "Terima kasih udh dengerin kita ngebac*t!"),
("Q&A","Silakan bertanya")
]

for title, content in slides_content:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = content

prs.save("Parallel_Processing.pptx")
print("PPT berhasil dibuat!")